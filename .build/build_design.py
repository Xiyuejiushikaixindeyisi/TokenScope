from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / '.build'
BUILD.mkdir(exist_ok=True)
OUT = ROOT / 'TokenScope_设计文档_V0.1.docx'
FONT = 'Microsoft YaHei'
NAVY = '213A51'
GRAY = '58636D'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.68)
section.bottom_margin = Inches(.62)
section.left_margin = Inches(.78)
section.right_margin = Inches(.78)
section.header_distance = Inches(.28)
section.footer_distance = Inches(.28)
WIDTH = 6.94

def font_style(style, size, bold=False, color='000000'):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), FONT)
    style.paragraph_format.widow_control = True

font_style(doc.styles['Normal'], 11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
font_style(doc.styles['Title'], 26, True)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
font_style(doc.styles['Subtitle'], 11)
font_style(doc.styles['Heading 1'], 18, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
font_style(doc.styles['Heading 2'], 12.5, True)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
font_style(doc.styles['Caption'], 9, color=GRAY)
font_style(doc.styles['List Bullet'], 11)
for s in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    ppr = doc.styles[s].element.find(qn('w:pPr'))
    if ppr is not None:
        for el in list(ppr):
            if el.tag in (qn('w:pBdr'), qn('w:shd')):
                ppr.remove(el)

hp = section.header.paragraphs[0]
hp.text = 'TokenScope  词元洞察'
hp.runs[0].font.size = Pt(9)
hp.runs[0].font.color.rgb = RGBColor(0, 0, 0)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run('讨论稿 V0.1  ·  ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
fp._p.append(field)
doc.core_properties.title = 'TokenScope 第一版设计文档'
doc.core_properties.subject = '面向企业 MaaS 的 Token 运营与自进化诊断设计'
doc.core_properties.author = 'TokenScope 项目组'
doc.core_properties.keywords = 'TokenScope, PrefixLens, MaaS, Token运营, SLO, Agent诊断'

def p(text='', bold_lead=None, style=None):
    x = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        x.add_run(bold_lead).bold = True
        x.add_run(text[len(bold_lead):])
    else:
        x.add_run(text)
    return x

def h(text, level=1):
    return doc.add_heading(text, level=level)

def newpage(title):
    heading = h(title)
    heading.paragraph_format.page_break_before = True

def bullets(items):
    for item in items:
        x = p(item, style='List Bullet')
        x.paragraph_format.space_after = Pt(5)

def table(headers, rows, widths, center_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9')
        borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'75' if side in ['top','bottom'] else '110'); e.set(qn('w:type'),'dxa'); margin.append(e)
            tcpr.append(margin)
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), NAVY if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF')); tcpr.append(shd)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.line_spacing=1.12
                if ci in center_cols:
                    par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.name=FONT
                    run.font.size=Pt(10)
                    run.font.bold=(ri==0)
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    gap=doc.add_paragraph()
    gap.paragraph_format.space_after=Pt(1)
    gap.paragraph_format.space_before=Pt(0)
    gap.paragraph_format.line_spacing=Pt(3)
    gap.add_run().font.size=Pt(3)
    return t

def caption(text):
    return p(text, style='Caption')

def link_par(label, url, desc=''):
    par=doc.add_paragraph()
    par.paragraph_format.space_after=Pt(6)
    rel=par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyp=OxmlElement('w:hyperlink'); hyp.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),NAVY); rp.append(color)
    fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'),FONT); rp.append(fonts)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'19'); rp.append(sz)
    run.append(rp); tx=OxmlElement('w:t'); tx.text=label; run.append(tx); hyp.append(run); par._p.append(hyp)
    if desc:
        r=par.add_run('  '+desc); r.font.size=Pt(9.5)
    return par

# Engineering figures only: the boxes represent system components and flow.
regular=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',29)
bold=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc',31)
small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',24)
def box(draw, xy, title, sub='', fill='#F1F5F8'):
    draw.rounded_rectangle(xy, radius=10, fill=fill, outline='#AABAC8', width=2)
    x1,y1,x2,y2=xy
    lines=title.split('\n')
    total=len(lines)*43+(42 if sub else 0)
    y=(y1+y2-total)/2
    for line in lines:
        draw.text(((x1+x2)/2,y),line,font=bold,fill='#152D42',anchor='mt')
        y+=43
    if sub:
        draw.text(((x1+x2)/2,y+2),sub,font=small,fill='#4E6170',anchor='mt')
def arrow(draw, a,b):
    draw.line([a,b],fill='#607B90',width=4)
    angle=math.atan2(b[1]-a[1],b[0]-a[0])
    pts=[b,(b[0]-16*math.cos(angle-.45),b[1]-16*math.sin(angle-.45)),(b[0]-16*math.cos(angle+.45),b[1]-16*math.sin(angle+.45))]
    draw.polygon(pts,fill='#607B90')
def image_doc(path,width=WIDTH):
    par=doc.add_paragraph()
    par.paragraph_format.space_after=Pt(4)
    par.add_run().add_picture(str(path),width=Inches(width))
    inline=par.runs[0]._r.xpath('.//wp:docPr')[0]
    inline.set('descr',path.stem.replace('_',' '))

im=Image.new('RGB',(1800,180),'white'); dr=ImageDraw.Draw(im)
labels=[('业务总览','选择异常窗口'),('任务调用链','定位异常任务'),('请求详情','查看内容与计量'),('归因与建议','证据支持结论'),('效果复核','记录收益与反馈')]
for i,(a,b) in enumerate(labels):
    x=10+i*360; box(dr,(x,15,x+330,165),a,b)
    if i<4: arrow(dr,(x+333,90),(x+357,90))
flow=BUILD/'product_flow.png'; im.save(flow)

im=Image.new('RGB',(1800,480),'white'); dr=ImageDraw.Draw(im)
for xy,ttl,sub in [((20,20,350,160),'请求与运行时事件','appid · task · usage · SLO'),((430,20,790,160),'模型适配与编码','模板 · Token · 字段映射'),((880,20,1240,160),'窗口与任务特征','缓存 · 长度 · 调用行为'),((1330,20,1780,160),'异常检测与规则归因','PrefixLens + 新探测器'),((430,295,790,450),'数据与版本存储','原始事件 · 特征 · 配置'),((880,295,1240,450),'休眠诊断 Agent','未知异常 → 候选规则'),((1330,295,1780,450),'工作台与验证记录','证据 · 建议 · 业务反馈')]:
    box(dr,xy,ttl,sub)
for a,b in [((352,90),(427,90)),((792,90),(877,90)),((1242,90),(1327,90)),((610,162),(610,290)),((1555,162),(1555,290)),((1243,372),(1327,372))]: arrow(dr,a,b)
# Unexplained cases go to the Agent; validated candidates return to the detector library.
dr.line([(1345,162),(1345,200),(1120,200),(1120,290)],fill='#607B90',width=4)
arrow(dr,(1120,270),(1120,290))
dr.text((1230,167),'未解释异常',font=small,fill='#4E6170',anchor='mt')
dr.line([(1040,293),(1040,255),(1415,255),(1415,165)],fill='#607B90',width=4)
arrow(dr,(1415,185),(1415,165))
dr.text((1230,259),'验证后更新规则',font=small,fill='#4E6170',anchor='mt')
arch=BUILD/'system_architecture.png'; im.save(arch)

# Page 1
doc.add_paragraph('TokenScope 第一版设计文档',style='Title')
doc.add_paragraph('词元洞察  面向企业 MaaS 的 Token 运营与自进化诊断',style='Subtitle')
meta=p('版本 V0.1 讨论稿    日期 2026年10月22日    面向 MaaS 运营 推理研发 与业务接入同事')
for r in meta.runs: r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(GRAY)
p('我们计划在 PrefixLens 的前缀缓存分析能力上建设 TokenScope：按租户 appid 和模型周期性分析 Token 使用、上下文组成与任务调用行为，定位造成额外计算和 SLO 退化的原因，并将未知问题的诊断经验沉淀为可复用探测器。')
p('第一版交付形态是内部 Web 工作台、异步分析服务、探测器库和按需唤醒的诊断 Agent。常规分析由确定性代码完成；Agent 只处理解释不足的显著异常，并生成经过验证的候选规则。')
h('1 业务背景与建设目标',2)
p('我们的 MaaS 面向 To B 企业。租户选择模型，平台提供约定并发，并按模型与请求长度承诺 P90 TTFT 和 P90 TPOT，不限制 Token 总量。TTFT 是首 Token 等待时间，TPOT 是生成阶段平均每 Token 耗时；两者的统计定义以服务合同为准。')
p('因此，同样的并发承诺下，冗余上下文、无效调用和缓存复用下降都会增加资源压力。TokenScope 的运营目标是在业务效果和 SLO 约束下，降低服务资源消耗，并缩短问题发现、定位及复核所需的人力时间。')
table(['价值方向','如何产生价值','如何衡量'],[
('资源成本','识别重复 Prefill、上下文膨胀及无进展调用，协助业务或平台优化。','同类负载下的卡时、SLO 达标吞吐、可避免计算量。'),
('排障人力','自动发现异常、聚合典型样本、定位字段或调用步骤，减少逐条读日志。','单事件人工分钟数、定位周期、重复问题自动归因率。'),
('经验复用','将一次调查形成的规则用于后续窗口及其他适用业务。','跨业务有效复用次数、误报率、未知异常占比。')
],[1.05,3.15,2.74])
p('本版范围：先覆盖前缀复用、任务调用膨胀和上下文增长三类问题，贯通“异常发现—证据归因—建议—效果复核”。实时阻断、自动修改业务 Prompt 和平台计费系统不纳入第一版。',bold_lead='本版范围：')
caption('文档用途：统一产品理解、讨论一期范围与试点资源。计划与验收门槛为建议值，不构成已完成能力或排期承诺。')

# Page 2
newpage('2 产品形态与使用流程')
p('TokenScope 面向 MaaS 运营、推理引擎研发和业务 Agent 开发者提供同一个分析入口。运营同事先发现受影响业务，研发同事下钻到请求和运行时证据，业务同事据此调整 Prompt、工具返回或调用逻辑。')
image_doc(flow)
caption('图 1 拟议的工作台操作路径')
table(['页面','主要内容','用户采取的行动'],[
('业务总览','按 appid、模型、时间窗口展示请求量、长度分布、理想与实际命中率、实际新计算输入量和 SLO。','选择异常窗口，按资源与 SLO 影响排序。'),
('任务与会话','调用树或时间线、累计 Token、调用次数、重试、上下文增长和任务状态。','找出贡献最大的任务与异常阶段。'),
('请求详情','原始请求/响应、渲染后输入、Token 区间、缓存断点、工具定义及结果占比。','核对具体字段、调用步骤和正常对照。'),
('异常详情','原因、证据、影响范围、置信度、建议、未解释部分和前后对比。','形成可执行改进项，登记业务反馈。'),
('探测器与调查','已有规则版本、适用范围、回放结果、Agent 调查进度及候选规则。','查看规则是否可推广，复核误报与运行成本。')
],[1.08,3.85,2.01])
h('一个典型使用场景',2)
p('某 Agent 业务的输入长度 P90 明显增加，缓存命中率仍然较高。系统将增长定位到少数任务，并指出某工具的大段返回内容在多轮请求中重复携带。业务同事核对样本、正常对照和调用链后，调整工具返回或上下文保留方式，再比较长度、TPOT 与任务完成情况。此场景为产品流程示例，尚非已验证结果。')
h('每条结论需要可追溯的证据',2)
p('字段差异可以解释缓存断点，修改后的收益仍需重新计算；重复动作可以触发疑似循环，确认无效行为还需要任务进展证据。报告区分观测变化、原因证据和验证结果，无法确认时保留“待验证”。')

# Page 3
newpage('3 PrefixLens 已有成果与继承方式')
p('PrefixLens 已提供理想命中率计算、低命中原因探测和前端展示，是 TokenScope 的直接基础。完整代码位于华为实验机，本机保存基础版。第一阶段将以实验机版本和回归样本为迁移依据。')
h('已有能力',2)
p('Prompt 诊断链路覆盖请求解析、Chat Template 渲染、Tokenizer 编码、Block 切分、前缀匹配、理想命中率、首个 Miss Block 定位，以及典型请求和修改建议。现有材料还描述了模型级与租户级报告、LCP 分布、Reuse Time、CLI 和交互前端。[P1][P2]')
p('案例报告同时描述了流量回放、目标 TTFT 对应的最低命中需求，以及 MTP、CP 和 Hybrid Attention 适配。它们作为待迁移资产纳入盘点；具体支持范围需要和实验机的代码版本、配置及测试逐项对应。[P2]')
table(['业务','根因','已有材料中的改善结果'],[
('ICSL','随机 cch 字段位于公共 Prompt 前部，破坏后续前缀复用。','理想命中率 36.5% → 90.2%；\n实际命中率 1.5% → 70%以上；\nP90 TTFT 降低 35 秒以上。'),
('财经','动态时间戳位于公共 Prompt 前部。','理想命中率约 27% → 80%以上；\n实际命中率约 20% → 60%以上；\nP90 TTFT 降低至少 10 秒。'),
('诊断协作','由人工长 Prompt 比对转为字段、样本和建议驱动。','材料记录定位周期由一周以上缩短至半天；两项业务在一周内完成 Prompt 修改。')
],[.72,2.65,3.57])
caption('表 1 结果引用 PrefixLens-PPT 现有案例材料。财经采用 outline.md 的保守口径；HTML 中另有更高数值。正式收益核算须统一时间窗口、流量范围和部署配置。上述结果不等同于 TokenScope 已实现的收益。')
h('第一版如何复用',2)
bullets([
'保留 PrefixLens 作为缓存诊断子引擎，将输入、输出和模型配置封装为稳定接口，先建立迁移前后结果一致的回归集。',
'优先接入完整版本已有的断点证据、原因分类和前端下钻能力；根据性能剖析逐步迁移热点，不以全量重写作为一期前提。',
'新增任务索引、窗口基线、上下文组成分析、Loop 探测与休眠 Agent，让一次性离线分析变成周期性的运营流程。'
])

# Page 4
newpage('4 指标口径与异常判断')
p('统一使用三个分析单位：请求用于精确计量与断点分析，任务用于观察调用过程，appid 与模型的时间窗口用于运营监控。所有指标携带模型与模板版本、采集覆盖率、统计窗口和分析配置。')
table(['指标','一期定义与用途'],[
('完整输入与组成','以实际服务端渲染输入为准，统计长度分布，以及 system、tools、history、tool result 等组成。'),
('理想命中率','在明确的缓存域、窗口初始状态和 Block 规则下，统计理论可复用 Token 占比；保留 PrefixLens 原始 Block 指标供核对。'),
('实际缓存与新计算输入','区分本地缓存、外部 KV 和新计算 Prefill；有引擎计量时优先使用，只有 usage 时标注推算口径。'),
('输出与累计任务量','分别记录输入、输出及可得的 reasoning 量；按任务累计调用次数、重试次数和 Token，保留未结束任务。'),
('任务进展信号','相同调用签名、重复错误、工具结果变化、上下文新增量和任务完成状态；用于识别疑似无进展行为。'),
('服务质量与资源','按合同模型和长度档位统计 P90 TTFT、P90 TPOT、SLO 违约情况，并关联并发、排队、实例和资源记录。')
],[1.55,5.39])
p('未命中 Token 不等于无效 Token。简化条件下，实际新计算输入量可由完整输入量减去实际复用量得到；抢占重算、投机解码、KV 传输和 Decode 成本仍需分别观察。Token 数不能直接换算为统一的算力成本。[R2]')
h('建立可比较的基线',2)
p('默认按 appid、模型 revision 和业务场景建立基线，必要时区分 Agent、Prompt 与工具集版本，并控制工作日和时段差异。低样本分组回退到较粗粒度。发布变更和已确认异常窗口单独标记，避免把故障学习成新的正常水平。')
p('告警同时考虑偏离幅度、持续时间、最小样本量和影响规模。长度分布采用分位数与分布距离，累计消耗同时拆解任务数、每任务调用数及每次调用长度。所有候选阈值通过历史正常窗口和已知异常回放校准。[R5][R7]')
p('短输入意图识别的低命中率可能合理；高命中率也可能伴随历史膨胀。第一版既检测相对退化，也定期运行低成本结构规则识别长期低效。跨实例或窗口聚合使用原始计数和直方图，不直接平均命中率或 P90。')

# Page 5
newpage('5 第一版归因能力')
h('前缀复用下降',2)
p('理想命中率下降时，调用 PrefixLens 检查首个断点、动态字段、工具顺序、工具 Schema 和消息变化，展示请求对照与位置映射。理想值稳定而实际命中下降时，转向实例路由、缓存驻留、淘汰和运行时配置证据；一期给出相关线索，复杂运行时因果分析按数据完整性逐步扩展。')
h('任务调用膨胀与疑似循环',2)
p('对任务内按时间排序的调用序列生成行为特征，优先覆盖重复调用且结果不变、重复错误重试、A 与 B 之间往返、相同内容重复读取，以及多轮调用缺少新增观察等情况。归一化签名只用于检测，缓存命中计算始终使用真实输入。')
p('循环结论分级表达：可观测的重复行为、疑似无进展、经工具状态或任务结果确认的无效过程。代码修改后的重复测试等正常行为应作为反例。统计总消耗时包括失败尝试，避免只看成功任务掩盖浪费。')
h('上下文增长与工具贡献',2)
p('对完整渲染输入建立字段来源映射，定位工具定义、工具结果、RAG 内容和历史消息的增长贡献。首批规则包括工具定义膨胀、大结果重复携带、历史重复拼接和异常序列化。将增长事件关联到工具、Agent 版本及后续 TPOT 变化。')
p('字段长度不能简单由各字段独立编码后相加。BPE 边界、模板控制 Token 和模型渲染方式可能影响结果；一期需保存来源映射和无法精确分配的边界部分，并核对组成与完整输入总量。')
table(['统一归因输出','要求'],[
('原因与适用性','reason_code、探测器版本、模型与业务适用条件、置信度。'),
('证据与位置','request_id / task_id、正常对照、字段或 Token 区间、调用步骤及差异。'),
('影响范围','受影响任务与请求、相关 Token 量、未解释部分，避免跨原因重复累计。'),
('改进与验证','建议动作、语义风险、可计算的反事实结果、待线上验证的资源和 SLO 效果。')
],[1.4,5.54])
p('通用能力与业务规则分离：相同规则可以复用于不同 appid，业务特有的参数、允许重试次数和工具语义通过配置提供。')

# Page 6
newpage('6 数据架构与高性能实现')
p('已确认可获取完整请求/响应、usage、缓存命中、耗时，以及 session、task、trace 标识。第一步是验证这些字段的关联质量与含义，并补齐模型、模板和部署配置指纹。')
image_doc(arch)
caption('图 2 第一版系统组成示意  所有异步任务保留输入版本与分析配置以支持复跑')
h('数据接入和异步处理',2)
p('请求事件及时落盘，分析服务默认每 2 小时增量运行一次，周期可配置。按 appid 和模型分区调度，通过水位、request_id 去重和幂等写入处理补传与重跑。跨窗口任务延续累计，迟到数据触发受影响窗口修订。')
p('保存原始事件、模型配置、编码结果、任务与窗口特征、异常证据及探测器版本。仅保留指标无法支持新原因调查，需在租户访问权限内保留可回溯样本，并按平台要求设置保留期限。')
h('缓存域与窗口的一致性',2)
p('按 appid 隔离的理想上限与允许跨租户共享的缓存域分别统计。冷启动窗口分析、带历史预热的分析明确标注，不混为同一基线。编码可并行，涉及历史参考和任务顺序的归因必须保留时序；随机抽样不能破坏命中历史和调用链。')
h('性能路径与模型适配',2)
p('Rust 优先承载解析、编码、Hash/Trie 和批量统计热点。复用服务端已有 Token ID 或精确编码缓存，一次编码供多个探测器使用；采用有界任务队列和线程预算，避免外层并发与 Tokenizer 内部并行相互争抢。批量聚合可借鉴 Polars。[R6]')
p('DeepSeek 适配优先研究 deepseek-recipe 的特殊 Token 区间与批量编码实现。其分段边界来自渲染规则，不能任意切文本，也不能直接假定适用于 GLM 或 Qwen。[R1]')
p('目标模型为 DeepSeek V4 系列、GLM 5.2/5.3、Qwen 3.6/3.8，支持状态以部署版本和一致性测试为准。适配固定 Tokenizer、Chat Template、thinking/tool 模式、Block 与并行配置。')

# Page 7
newpage('7 休眠 Agent 与探测器自进化')
p('Agent 的职责是调查显著但解释不足的异常，并把有效调查沉淀为确定性探测器。常规窗口继续运行已有规则，避免对全部请求重复进行大模型分析。选择性介入可借鉴 SupervisorAgent，但离线规则生成和验证需由 TokenScope 自行建设。[R4]')
h('唤醒与调查',2)
p('仅在数据质量合格、异常影响达到门槛，并且已有探测器覆盖不足、置信度低或结论冲突时唤醒。相似事件先聚类去重，共享调查结果；同一异常设置冷却时间与单次分析预算。数据缺失和配置错配优先交给数据质量检查处理。')
p('Agent 获得异常窗口、正常对照、代表性任务、字段映射和已有规则结果，在只读样本范围内提出假设、编写启发式分析代码并生成证据。候选代码隔离运行，限制执行时间、内存和外部访问。无法确认的结论保留未解释状态。')
h('候选规则如何进入共享库',2)
table(['阶段','产出','进入下一阶段的条件'],[
('形成候选','原因定义、特征、检测代码、适用条件和运行成本。','可以在原始异常样本上复现，并提供具体证据。'),
('验证正确性','历史回放、正例、反例及独立留出样本。','核对误报、漏报与归因一致性；不能仅验证触发它的事件。'),
('验证可迁移性','其他业务的适用样本和正常样本测试。','通用逻辑与业务参数分离，确认适用边界。'),
('试运行与发布','影子运行记录、规则版本和回滚记录。','第一版经规则维护者复核后启用；观察反馈并可撤回。')
],[1.1,2.75,3.09])
p('跨业务推广复用规则和适用配置；原始租户内容保持在各自权限域内。第一版 Agent 可以自动分析和生成候选探测器，生产规则启用保留复核环节，业务 Prompt、Agent 代码和线上流量策略由原责任团队调整。')
p('每个事件记录采样、定位、沟通、规则复核和验证的人力时间。自进化的价值体现为同类问题自动归因，以及 A 业务的规则在 B 业务有效复用；净收益需扣除复核、误报处理与维护投入，并同时观察未知原因占比和误报率。')

# Page 8
newpage('8 第一版范围与实现计划')
p('第一版完成三类诊断主线、内部工作台和受控的探测器自进化闭环。先选择 Agent 编程、短输入意图识别、长上下文或 RAG 三类代表业务；模型接入按试点实际流量排序，后续使用同一适配验收流程扩展。')
table(['阶段','建议时间','主要交付','阶段验收'],[
('P0 基础对齐','第 1 周','盘点实验机 PrefixLens；确认字段、SLO、模型配置与案例；形成回归样本。','可复现既有案例，明确指标和支持边界。'),
('P1 数据与总览','第 2 至 3 周','增量采集、任务索引、模型适配、窗口特征和业务总览。','可追溯至请求；批次可重跑；核心计量对账通过。'),
('P2 三类诊断','第 4 至 5 周','接入 PrefixLens；实现上下文组成和首批 Loop 规则；完成异常下钻。','三类业务各有正例和正常反例，输出可解释证据。'),
('P3 自进化试运行','第 6 至 7 周','异常聚类、Agent 按需唤醒、候选生成、隔离回放及规则版本管理。','至少完成一次未知问题到可复用规则的演练。'),
('P4 试点与交付','第 8 周','性能优化、连续批次运行、业务反馈、收益对照和使用文档。','按下页门槛评审，形成下一阶段推广清单。')
],[.94,.88,2.92,2.2])
caption('表 2 以上为约 8 周的初步估算，假设 2 名研发持续投入，并有推理/数据接口和试点业务同事配合。实验机资产盘点、数据规模和资源落实后重新评估。')
h('建议的协作分工',2)
bullets([
'MaaS 运营：确定试点业务、异常优先级、SLO 档位和收益评价方式。',
'平台与推理研发：对齐请求级计量、模型配置、运行时指标和服务端结果。',
'TokenScope 研发：交付采集分析、归因引擎、工作台、性能优化与规则管理。',
'业务 Agent 开发者：提供版本和任务语义，确认异常，实施改进并验证业务效果。'
])
h('第一版边界',2)
p('纳入：周期批量分析、三类异常、证据下钻、建议与复核记录、候选规则自动生成和验证。后续扩展：在线阻断、无人值守启用规则、自动修改业务代码、精确到任务的硬件成本归因，以及覆盖全部模型与全部租户的常态化运营。')

# Page 9
newpage('9 验收标准与收益验证')
p('验收分为计量正确、诊断可信、运行可持续和业务价值四部分。第一版优先证明数据与证据闭环成立；资源节省和人力降低幅度通过试点形成基线，不预先承诺统一百分比。')
table(['验收项','第一版建议门槛'],[
('输入编码和计量','对每个已接入模型与配置建立固定样本集，Token ID 与服务端一致；usage、缓存和新计算量差异能够逐项解释。'),
('数据完整与任务关联','建议有效事件关键字段完整率及可关联率均不低于 99%；缺失、迟到、重复和未结束任务有独立统计。'),
('三类诊断能力','每类都有已知异常与正常反例；分原因报告人工确认准确率、漏检、误报和未解释比例。建议试点告警确认准确率达到 90%，低样本原因保留人工复核。'),
('分析吞吐与稳定性','在约定试点峰值规模下，建议 2 小时窗口的批次分析不超过 1 小时，连续运行 7 天无持续积压；报告 CPU、内存、编码吞吐和缓存复用。'),
('自进化闭环','至少一个候选规则通过独立样本、跨业务适用性验证和影子运行；能够追溯版本并回滚。'),
('产品可用性','从业务窗口定位到任务、请求和证据；可记录处理人、建议动作、业务反馈及效果对照。')
],[1.32,5.62])
caption('表 3 数值为试点建议门槛，需在 P0 按数据规模与可用样本确认；样本数、评估范围和未覆盖原因应随验收结果一起报告。')
h('资源成本如何核算',2)
p('固定模型、硬件与服务配置，以同一任务集或业务组成、任务难度和并发可比的流量做回放或受控灰度。输入长度和调用次数允许随优化变化。先验证任务质量与 SLO，再测量单位成功任务资源消耗或 SLO 达标吞吐；缺少任务结果时，先报告请求级资源和性能。')
p('以实际减少的卡时乘内部单位成本，形成可核算资源节省；仅释放容量但未缩容时，报告容量收益或延后扩容空间。人力收益按节约小时与内部成本核算，再扣除 TokenScope 分析、Agent 推理、误报处理及维护成本，避免重复计算。')
h('试点需要形成的决策',2)
p('试点结束后确认：三类问题的真实发生频率、哪些探测器值得推广、哪些模型配置仍需适配、批次周期是否足够，以及自动化节省是否覆盖运行和维护投入。以这些结果决定第二阶段扩大覆盖面或增加在线干预。')

# Page 10
newpage('附录 依据与参考实现')
p('已有成果采用 PrefixLens-PPT 案例材料，产品范围采用本次业务讨论中确认的服务模式和数据条件。公开项目用于实现参考，论文用于设计依据，不将其基准结果作为 TokenScope 的收益承诺。')
h('内部依据',2)
p('[P1] PrefixLens-ppt/PrefixLens/outline.md：已有诊断能力、ICSL 与财经结果，以及诊断周期。财经案例使用其中的保守值。')
p('[P2] PrefixLens-ppt/PrefixLens/prefixlens_report.html：完整版本产品描述、前缀诊断与流量回放、运行时适配说明。材料中的界面占位图不作为真实产品截图引用。')
p('[P3] PrefixLens/README.md、USAGE.md 与本机核心源码：基础版输入编码、理想命中统计、模型/APP 报告和 Prompt 变体分析。完整能力清单以实验机盘点为准。')
h('公开代码与规范',2)
link_par('[R1] DeepSeek Recipe', 'https://github.com/deepseek-ai/deepseek-recipe', 'Rust 渲染、特殊 Token 区间和批量编码；重点文件 deepseek-recipe-encoding/src/tokenizer.rs。')
link_par('[R2] vLLM 请求与 Prefill 统计', 'https://github.com/vllm-project/vllm/blob/main/vllm/v1/metrics/stats.py', '本地计算、本地缓存、外部 KV 与请求统计的数据结构。')
link_par('[R3] OpenTelemetry GenAI 语义约定', 'https://github.com/open-telemetry/semantic-conventions-genai', '用于模型、Agent、工具调用和 usage 字段的接入参考，采用时固定版本。')
link_par('[R4] SupervisorAgent', 'https://github.com/LINs-lab/SupervisorAgent', '低成本过滤后选择性调用监督 Agent；用于触发与调查流程参考。')
link_par('[R5] Evidently 数据漂移', 'https://docs.evidentlyai.com/metrics/explainer_drift', '参考批次与当前批次的分布比较，阈值需按实际业务校准。')
link_par('[R6] Polars 流式处理', 'https://docs.pola.rs/user-guide/concepts/streaming/', '用于大规模离线聚合与内存控制。')
h('相关论文',2)
link_par('[R7] How Do AI Agents Spend Your Money', 'https://arxiv.org/abs/2604.22750', '2026；编程 Agent 的 Token 波动、输入占比与任务阶段分析。')
link_par('[R8] Stop Wasting Your Tokens', 'https://arxiv.org/abs/2510.26585', 'ICLR 2026；SupervisorAgent 的选择性监督方法。')
link_par('[R9] Early Diagnosis of Wasted Computation', 'https://arxiv.org/abs/2606.01365', '2026 预印本；重复行为、工具故障、低信息增益与选择性深入评估。')
link_par('[R10] Why Do Multi Agent LLM Systems Fail', 'https://arxiv.org/abs/2503.13657', '2025；MAST 失败分类和标注轨迹，可作为原因体系与验证集参考。')

doc.save(OUT)
print(str(OUT))
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'inline_shapes',len(doc.inline_shapes))
